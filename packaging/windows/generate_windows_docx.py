"""Regenerate MeetingBox Windows.docx from WINDOWS_DESKTOP.md-style content."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
OUT_REPO = ROOT / "MeetingBox_Windows.docx"
OUT_DESKTOP = Path(r"C:\Users\shivakumar\Desktop\MeetingBox Windows.docx")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_table_row(table, col0: str, col1: str) -> None:
    row = table.add_row().cells
    row[0].text = col0
    row[1].text = col1


def build() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("MeetingBox Windows — Desktop Release Guide", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    doc.add_paragraph(
        "Scope: the Windows desktop port (companion + dashboard installer). "
        "The Linux appliance is out of scope for this document."
    )
    doc.add_paragraph(
        "Last updated: July 2026. Source of truth in the repo: "
        "mini-pc/packaging/windows/WINDOWS_DESKTOP.md and BUILD.md."
    )

    add_heading(doc, "What ships today", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Capability"
    hdr[1].text = "Status"

    rows = [
        ("Installer", "MeetingBoxSetup.exe — one installer, two apps (companion + dashboard)"),
        ("Companion", "MeetingBox.exe + meetingbox-audio.exe (PyInstaller); audio child auto-managed"),
        ("Dashboard", "MeetingBoxDashboard.exe (Tauri + WebView2 + React SPA)"),
        ("Windows backend", "https://win.meetingboxai.lucratechsol.com (Linux appliance unchanged)"),
        ("Desktop UX", "No Settings screen; no system-check overlay on Windows"),
        ("Legal", "Installer EULA (EULA.rtf); dashboard first-login consent gate"),
        ("Compliance file", "THIRD-PARTY-NOTICES.txt shipped with install"),
        ("WebView2", "Bundled bootstrapper; silent install if runtime missing"),
        ("Auto-start", "HKLM Run for both apps (opt-out install task); dashboard uses --minimized"),
        ("Theme", "Dashboard light + periwinkle (#7b61ff); matches companion / brand site"),
        ("Companion window", "Fixed 1120×680, clamped to work area (fits 1366×768 laptops)"),
        ("Upgrade", "Same AppId — in-place upgrade; CloseApplications=force during install"),
        ("Uninstall", "taskkill companion, audio child, dashboard before file removal"),
        ("Signing", "sign.ps1 + Inno /DSIGN wired; EV/OV cert required for production download"),
    ]
    for a, b in rows:
        add_table_row(table, a, b)

    add_heading(doc, "Install layout", 1)
    for line in [
        "%ProgramFiles%\\MeetingBox\\ — MeetingBox.exe, meetingbox-audio.exe, _internal\\, Dashboard\\",
        "%ProgramData%\\MeetingBox\\device-ui.env — seeded first install only; kept on uninstall",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "SmartScreen (unsigned builds)", 1)
    doc.add_paragraph(
        "Local or CI builds compiled without /DSIGN show “Windows protected your PC”. "
        "For testing: More info → Run anyway. For customers: sign all three exes and the "
        "installer with an EV (recommended) or OV certificate."
    )

    add_heading(doc, "Upgrade and uninstall", 1)
    for line in [
        "Re-running MeetingBoxSetup.exe upgrades the same product; it does not install a duplicate entry.",
        "Running apps are closed automatically during upgrade (no files-in-use prompt).",
        "Uninstall from Control Panel stops all MeetingBox processes, then removes Program Files files.",
        "Older installers (before the uninstall fix) may leave processes running — end tasks once, then use the latest installer.",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "Auth note", 1)
    doc.add_paragraph(
        "Companion uses the device pairing token; dashboard uses user JWT in WebView2 storage. "
        "Sign-in is not automatically synchronized between the two apps today."
    )

    add_heading(doc, "Build order (summary)", 1)
    for i, line in enumerate(
        [
            "PyInstaller → packaging\\windows\\dist\\MeetingBox\\",
            "frontend: npm run tauri:build → MeetingBoxDashboard.exe",
            "Stage MicrosoftEdgeWebview2Setup.exe beside MeetingBox.iss",
            "(Release) sign.ps1, then ISCC with /DSIGN",
            "Output: packaging\\windows\\Output\\MeetingBoxSetup.exe",
        ],
        start=1,
    ):
        doc.add_paragraph(f"{i}. {line}")

    add_heading(doc, "QA checklist (clean VM)", 1)
    for line in [
        "Signed installer: publisher shown; no SmartScreen block (EV)",
        "EULA + dashboard consent gate",
        "WebView2 on VM without runtime",
        "Auto-start and tray behavior",
        "Companion window fits small screens",
        "win.meetingboxai backend: OAuth redirect URI, voice/recording",
        "Upgrade while apps running; clean uninstall (no processes in Task Manager)",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "Roadmap (not yet in product)", 1)
    for line in [
        "Google OAuth verification + CASA for restricted scopes",
        "Licensing, subscriptions, payments",
        "Desktop auto-update",
        "Crash reporting / analytics (opt-in)",
        "Legal review of EULA and privacy copy",
        "Optional MSIX / Microsoft Store",
    ]:
        add_bullet(doc, line)

    add_heading(doc, "Commercial checklist (planning)", 1)
    doc.add_paragraph(
        "The sections below remain a planning backlog for a full commercial launch. "
        "Items marked “wired” or “shipped” above reflect what is already in the Windows build."
    )

    sections = [
        (
            "Legal and business",
            [
                "EULA in installer — shipped (counsel review recommended)",
                "Privacy policy / ToS / DPA for cloud — publish on web; link from consent gate",
                "Recording consent — in-app gate shipped; meeting-room guidance for end users",
                "THIRD-PARTY-NOTICES — shipped",
            ],
        ),
        (
            "Platform compliance",
            [
                "Google OAuth verification — required for production Gmail/Calendar; register win.* redirect URI",
                "OpenAI commercial terms and user disclosure",
                "GDPR/CCPA flows on backend",
            ],
        ),
        (
            "Trust and distribution",
            [
                "EV code signing — build wired; acquire cert and run sign.ps1 + /DSIGN",
                "Marketing site download with system requirements",
            ],
        ),
        (
            "Product gaps",
            [
                "License enforcement and payments",
                "Auto-update for desktop",
                "Telemetry / crash reporting",
            ],
        ),
    ]
    for heading, bullets in sections:
        add_heading(doc, heading, 2)
        for b in bullets:
            add_bullet(doc, b)

    return doc


def main() -> None:
    doc = build()
    doc.save(OUT_REPO)
    try:
        doc.save(OUT_DESKTOP)
        print(f"Wrote {OUT_REPO}")
        print(f"Wrote {OUT_DESKTOP}")
    except OSError as e:
        print(f"Wrote {OUT_REPO}")
        print(f"Could not write Desktop copy: {e}")


if __name__ == "__main__":
    main()
